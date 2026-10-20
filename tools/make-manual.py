# -*- coding: utf-8 -*-
"""Genera el manual de usuario de IntegrityWatch Field en DOCX (python-docx)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EV = os.path.join(ROOT, 'docs', 'evidence')
OUTDIR = os.path.join(ROOT, 'docs', 'manual')
OUT = os.path.join(OUTDIR, 'IntegrityWatch-Field-Manual-Usuario.docx')
os.makedirs(OUTDIR, exist_ok=True)

doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)

ACCENT = RGBColor(0x0B, 0x5C, 0xAD)
MUTED = RGBColor(0x5A, 0x6B, 0x7C)


def h1(t):
    return doc.add_heading(t, level=1)


def h2(t):
    return doc.add_heading(t, level=2)


def p(t, bold=False, italic=False):
    para = doc.add_paragraph()
    run = para.add_run(t)
    run.bold = bold
    run.italic = italic
    return para


def steps(items):
    for i, s in enumerate(items, 1):
        doc.add_paragraph(f"{i}. {s}")


def bullets(items):
    for s in items:
        doc.add_paragraph(s, style='List Bullet')


def table(header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    for i, hcell in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = hcell
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()


def img(name, caption):
    path = os.path.join(EV, name)
    if not os.path.exists(path):
        doc.add_paragraph(f"[imagen no encontrada: {name}]")
        return
    doc.add_picture(path, width=Cm(16.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED


# ---------- PORTADA ----------
title = doc.add_paragraph()
tr = title.add_run('IntegrityWatch Field')
tr.bold = True
tr.font.size = Pt(28)
tr.font.color.rgb = ACCENT
sub = doc.add_paragraph('Manual de usuario — Centro de operaciones de integridad y cumplimiento, offline-first, para trabajo en campo petrolero.')
sub.runs[0].font.size = Pt(13)
doc.add_paragraph('Versión 1.0 · Documento generado el 2026-10-19 · Aplicación funcional sin conexión (PWA + escritorio).')
warn = doc.add_paragraph('Herramienta de apoyo. Valida siempre los resultados contra las ediciones vigentes de las normas y los datos del fabricante de los equipos.')
for r in warn.runs:
    r.bold = True
doc.add_paragraph()

# ---------- 1 INSTALACIÓN ----------
h1('1. Instalación y arranque')
p('La aplicación funciona completamente en tu equipo, sin internet y sin servidor obligatorio. Hay tres formas de arrancarla:')
table(
    ['Vía', 'Pasos', 'Cuándo usarla'],
    [
        ['A. START-HERE.bat (recomendada)', 'Doble clic en START-HERE.bat → se abre el navegador en http://127.0.0.1:8419', 'Uso diario; permite instalar la app como PWA'],
        ['B. Doble clic en index.html', 'Abre el archivo directamente (modo file://)', 'Prueba rápida; sin service worker ni PWA'],
        ['C. Servidor estático propio', 'tools\\serve.ps1 o tools\\serve.py (o nginx/IIS con los tipos MIME correctos)', 'Equipos con políticas o servidores existentes'],
    ],
)
p('Instalación como PWA (Chrome/Edge): con la app abierta por START-HERE, pulsa el icono de instalación en la barra de direcciones (o menú ⋮ → “Instalar aplicación”).')
p('Ejecutable de escritorio: app\\IntegrityWatch Field 1.0.0.exe (portable, no requiere instalación). Nota: al no estar firmado digitalmente, Windows SmartScreen puede pedir “Más información → Ejecutar de todas formas”.')
p('Verificación del paquete portable: la carpeta incluye SHA256SUMS.txt con el hash SHA-256 de cada archivo. Tras descargar o copiar el paquete, recalcula los hashes (certutil -hashfile archivo SHA256 en Windows) y compáralos: si coinciden, el paquete llegó intacto.', bold=True)

# ---------- 2 PRIMEROS PASOS ----------
h1('2. Primeros pasos')
steps([
    'Al abrir por primera vez se cargan datos de demostración (puntos de corrosión, válvulas, ductos, activos APM y lecturas) para que veas todo funcionando.',
    'Elige idioma en el encabezado: English, Español, Français, 简体中文, Русский, العربية (interfaz de derecha a izquierda) o Português.',
    'Elige tema con el botón del encabezado: claro → oscuro → automático (se recuerda solo).',
    'El selector de sector (Todos / Upstream / Midstream / Downstream) filtra el Panel y el APM.',
    'Para empezar de cero: Admin → Restablecer datos de demo (o borra todos los datos).',
])
img('p0-locale-es.png', 'Panel principal en español: KPIs con contexto, mapa de calor por severidad y navegación agrupada por función.')

# ---------- 3 PANEL ----------
h1('3. Panel de operaciones')
p('La pantalla inicial resume el estado de integridad de todo el campo:')
table(
    ['KPI', 'Qué muestra', 'Referencia en pantalla'],
    [
        ['Tasa de corrosión promedio', 'Media de mpy de los puntos de monitoreo, con gauge analógico', 'Aceptable < 5 mpy'],
        ['Puntos severos', 'Puntos con lectura ≥ 10 mpy (SP0775)', 'Umbral SP0775: ≥ 10 mpy'],
        ['PSV vencidas', 'Válvulas de alivio con prueba vencida', 'n de N válvulas'],
        ['Alertas APM', 'Activos con recomendaciones abiertas', 'n de N activos'],
        ['Cumplimiento pendiente', 'Recomendaciones abiertas', 'n de N abiertas'],
    ],
)
p('El mapa de calor muestra una celda por punto de monitoreo: el color y la forma indican la severidad (● Baja, ■ Moderada, ▲ Alta, ✕ Severa). La fuente de alertas lista las recomendaciones abiertas ordenadas por gravedad. El banner rojo/ámbar bajo el encabezado aparece cuando hay alarmas activas (ASM/API-1165).')
p('El botón “Ejecutar reglas” analiza todos los datos y genera recomendaciones; también corre automáticamente cada 60 segundos (configurable en Admin).')
img('01-dashboard-es-light.png', 'Panel con KPIs, mapa de calor y fuente de alertas.')

# ---------- 4 CORROSIÓN ----------
h1('4. Corrosión')
h2('4.1 Crear un punto de monitoreo')
steps([
    'Corrosión → botón “Nuevo punto”.',
    'Completa: código (p. ej. MP-CORR-021), nombre, sector, sistema de ducto (opcional), línea, material (API 5L Gr B, X52, X60, X70, A106B, 316L, 317L, dúplex 2205, superdúplex 2507, acero con recubrimiento…), tipo de sensor (cupón gravimétrico, sonda ER, LPR, H₂, UT, visual).',
    'Parámetros de proceso: H₂S, CO₂, presión total, temperatura, BS&W, velocidad, objetivo y residual de inhibidor.',
    'Protección catódica y aislamiento (CUI) si aplican. Guarda.',
])
h2('4.2 Lecturas y trend chart')
steps([
    'Entra al punto → “Nueva lectura” → fecha + tasa (mpy) + nota.',
    'El TREND CHART muestra la serie con selector de rango (24 h · 7 d · 30 d · 90 d · Todo), líneas de umbral SP0775 (5 mpy aceptable, 10 mpy severo) y cursor con fecha, valor y clasificación (mueve el mouse o usa las flechas del teclado).',
])
h2('4.3 Cupones y eficiencia de inhibidor')
steps([
    '“Nuevo cupón”: Rol (Blanco/Inhibido), Par (p. ej. P3), peso inicial/final, área, días. La tasa se calcula sola (ASTM G4: CR = 534·ΔW/(ρ·A·t)).',
    'El panel “Eficiencia del inhibidor” compara el cupón blanco y el inhibido del mismo par: EI = (CRb − CRi)/CRb × 100.',
    'Puedes editar cualquier cupón con el botón Editar de su fila (rol, par, pesos, material, fechas).',
])
h2('4.4 Paneles automáticos del punto')
bullets([
    'Balance de inhibidor: residual vs objetivo (ventana 80–150 %).',
    'De Waard–Milliams: modelo dulce vs medido (ratio).',
    'Screening agrio MR0175: pH₂S = P × fracción molar.',
    'Protección catódica: criterio −850 mV vs CSE o 100 mV de despolarización.',
    'CUI: aislamiento y última inspección.',
])
p('“Expediente del punto” genera el expediente normativo completo listo para imprimir o guardar como PDF.')
img('og-trend-chart.png', 'Trend chart con umbrales SP0775, selector de rango y tooltip de clasificación.')
img('v2-03-corrosion-detail-light.png', 'Ficha del punto con paneles de cálculo.')

# ---------- 5 VÁLVULAS ----------
h1('5. Válvulas de alivio (PSV)')
h2('5.1 Registrar una válvula')
steps([
    'Válvulas de alivio → “Nueva PSV”.',
    'Completa: etiqueta, tipo (convencional / fuelle compensado / pilotada), fluido (11 pre-cargados, editables), presión de ajuste, temperatura, contrapresión, fechas de prueba.',
])
h2('5.2 Dimensionar (API 520 Parte I)')
steps([
    'En la ficha de la válvula, pestaña Gas: caudal W, P1 (o ajuste + sobrepresión), T, Z, M, k, Kd, Kb/Kc → área requerida; el sistema detecta flujo crítico o subcrítico (F2).',
    'Pestaña Líquido: caudal Q, gravedad G, ΔP, Kd/Kw/Kc y viscosidad → área con corrección Kv iterativa (Re en el orificio).',
    'El sistema selecciona el orificio API 526 (D–T), calcula margen y capacidad real. “Guardar dimensionamiento” lo asocia a la válvula.',
])
p('“Datasheet” genera la ficha API 520/ISA S20; puedes copiarla al portapapeles o imprimirla a PDF.')
img('03-psv-detail-sizing.png', 'Ficha de válvula con panel de dimensionamiento.')

# ---------- 6 DUCTOS ----------
h1('6. Ductos (midstream)')
steps([
    '“Nuevo sistema”: nombre, operador, producto, diámetro, espesor, material, recubrimiento, longitud y rango de estaciones (km).',
    'Dentro del sistema: “Nuevo activo de línea” → tipo (12 tipos: válvulas de seccionamiento/bloqueo, bridas, estaciones de bombeo/compresión, tanques, sumideros, cámaras, cruces, ROW, medición), estación km y lado.',
    '“Nueva inspección”: tipo (ILD, corrosión externa, visual, UT, reparación, reemplazo), tramo km, severidad, hallazgos.',
    'El MAPA dibuja los activos a escala de km con color y forma por severidad de la última inspección; navegable con Tab y Enter; filtros por tipo y severidad.',
])
img('v2-04-map-light.png', 'Mapa esquemático del poliducto con activos por estación km.')

# ---------- 7 APM ----------
h1('7. Desempeño de activos (APM)')
h2('7.1 Crear un activo')
steps([
    'Desempeño de activos → “Nuevo activo”.',
    'Código, nombre, sector y categoría (incluye equipos rotativos y de seguridad: bomba centrífuga API 610, bomba de desplazamiento positivo, generador ISO 8528, subestación, bomba contra incendio NFPA 20, sistema de extinción…).',
    'Criticidad A/B/C, estado, POF y COF (1–5) → la banda de riesgo se calcula sola (API 580/581); intervalo base de inspección y fecha siguiente sugerida.',
])
h2('7.2 Eventos y métricas')
steps([
    '“Nuevo evento”: fallo, mantenimiento, parada no planificada, orden de trabajo o inspección, con fechas y tiempo de parada.',
    'El sistema calcula disponibilidad, MTBF, MTTR y el semáforo de salud; la matriz de riesgo 5×5 muestra la prioridad de inspección.',
])
img('og-apm-equipment.png', 'Lista de activos con equipos rotativos y de seguridad añadidos.')

# ---------- 8 INGESTA ----------
h1('8. Ingesta de datos')
h2('8.1 CSV')
steps([
    'Ingesta de datos → pestaña CSV → “Elegir archivo” (o arrastra el archivo). Acepta , ; y tabulador.',
    'Revisa el mapeo de columnas (automático; admite variantes con tildes en español/inglés).',
    'Marca “Crear puntos faltantes” si el archivo trae puntos nuevos, y vuelve a ejecutar “Ejecutar validación”.',
    'La tabla de validación muestra los errores por fila (fecha futura, Wf > Wi, fuera de rango, punto inexistente) y omite duplicados automáticamente.',
    '“Importar” guarda lo válido y registra la operación en el historial.',
])
h2('8.2 WITSML-lite y JSON')
p('La pestaña WITSML importa pozos, trayectorias y canales simplificados. La pestaña JSON exporta/importa respaldos completos (modo merge o reemplazo) — úsala para respaldos semanales.')
h2('8.3 Cola offline')
p('Si configuras un endpoint en Admin, los cambios se encolan en IndexedDB y se envían al recuperar conexión. Sin endpoint, todo queda local.')

# ---------- 9 RECOMENDACIONES ----------
h1('9. Recomendaciones')
steps([
    'Filtra por estado (Abierta/Reconocida/Resuelta) y gravedad.',
    'Cada tarjeta muestra la regla, el mensaje con datos reales, la entidad afectada y el enlace a su ficha.',
    'Acciones: Reconocer → Resolver (con nota). Todo queda auditado con fecha y usuario.',
    'Si la condición vuelve a cumplirse tras resolver, la recomendación se reabre automáticamente.',
    'El banner rojo/ámbar bajo el encabezado siempre muestra las alarmas activas.',
])

# ---------- 10 INFORMES ----------
h1('10. Informes')
steps([
    'Informe trimestral: Informes → Informe trimestral → elige fechas → Generar. Incluye portada, KPIs, gráfico semanal, top de puntos críticos, eficiencia por par, campañas y recomendaciones.',
    'Expediente por punto y datasheet de válvula: mismos botones.',
    'Pulsa Imprimir → “Guardar como PDF” en el diálogo del navegador.',
])
img('06-quarterly-report.png', 'Informe trimestral listo para imprimir o guardar como PDF.')

# ---------- 11 GLOSARIO ----------
h1('11. Glosario')
p('Buscador y filtro por categoría (corrosión, PSV, ductos, APM, normas, datos) con 65 términos técnicos y su referencia normativa, en los 7 idiomas de la aplicación.')
img('07-glossary-en.png', 'Glosario técnico multilingüe.')

# ---------- 12 ADMIN ----------
h1('12. Administración')
table(
    ['Sección', 'Para qué'],
    [
        ['Apariencia', 'Tema e idioma (equivalente al encabezado)'],
        ['Reglas', 'Intervalo de barrido y umbrales (banda residual, disponibilidad, sensor atascado, z-score)'],
        ['Datos', 'Exportar todo a JSON, importar respaldo, restablecer demo, borrar todo'],
        ['Seguridad', 'Bloqueo con PIN (roles admin/operador)'],
        ['Almacenamiento', 'Uso de espacio y conteo de registros'],
        ['Aplicación', 'Forzar actualización: limpia la caché del service worker y recarga la última versión publicada'],
    ],
)

# ---------- 13 PROBLEMAS ----------
h1('13. Solución de problemas')
table(
    ['Síntoma', 'Causa probable', 'Solución'],
    [
        ['El puerto está ocupado', 'Otro proceso usa 8419', 'tools\\serve.ps1 -Port 9000 (o cierra el otro proceso)'],
        ['PowerShell bloquea el script', 'Política de ejecución', 'START-HERE.bat ya incluye el bypass; alternativa: py tools\\serve.py'],
        ['No veo la opción de instalar PWA', 'Abriste por file://', 'Usa START-HERE.bat'],
        ['Sigo viendo la versión anterior', 'El service worker conserva la versión previa (cache-first)', 'Pulsa “Recargar” en el aviso; si persiste: Ctrl+F5 o Admin → Forzar actualización'],
        ['Los datos desaparecieron', 'Modo incógnito o borrado de datos del sitio', 'Restaura tu respaldo JSON; evita el modo privado'],
        ['El informe sale en blanco', 'Se generó y cambió de vista', 'Vuelve a Generar y pulsa Imprimir sin salir'],
        ['Errores de validación al importar', 'Filas con fecha/Wf/rango fuera de norma', 'Corrige las filas señaladas y vuelve a importar'],
        ['No aparece mi punto en el mapa', 'El activo está fuera del rango km del sistema', 'Ajusta stationKm del activo o el rango del sistema'],
        ['SmartScreen bloquea el EXE', 'Binario sin firma digital', 'Más información → Ejecutar de todas formas'],
        ['La app está lenta con muchos datos', 'Volumen histórico grande', 'Filtra por sector/periodo; el soporte de paginación está en el roadmap'],
    ],
)

# ---------- 14 BUENAS PRÁCTICAS ----------
h1('14. Buenas prácticas de campo')
bullets([
    'Respalda cada semana (Admin → Exportar JSON) y antes de importaciones grandes.',
    'Valida antes de importar; revisa siempre la tabla de errores.',
    'Usa códigos consistentes (MP-CORR-nnn, PSV-nnn, SYS-…) para que los informes queden legibles.',
    'Activa el PIN si el equipo es compartido.',
    'Los cálculos son de apoyo: valida siempre contra las ediciones vigentes de las normas y los datos del fabricante.',
])

doc.save(OUT)
size_kb = round(os.path.getsize(OUT) / 1024, 1)
doc2 = Document(OUT)
paras = len(doc2.paragraphs)
tables = len(doc2.tables)
images = sum(1 for rel in doc2.part.rels.values() if 'image' in rel.reltype)
print(f"OK {OUT}")
print(f"KB={size_kb} parrafos={paras} tablas={tables} imagenes={images}")
